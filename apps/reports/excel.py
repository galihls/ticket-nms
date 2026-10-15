"""Builder Excel laporan harian & bulanan.

Format mengacu pada `Report Monitoring INITECH 2026.xlsx`. Selama file acuan
belum diterima, struktur sheet di sini adalah pendekatan terbaik dari
kebutuhan: satu sheet rekap, satu sheet detail tiket, satu sheet performa
agent. Kolomnya tinggal digeser kalau file acuan sudah ada.
"""

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY = "16294F"
NAVY_SOFT = "EDF1F8"
LINE = "E3E8EF"

STATUS_FILL = {
    "Open": "EBF2FE",
    "Outstanding": "FEF3E8",
    "Close": "E7F6EE",
}

_thin = Side(style="thin", color=LINE)
BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _title_block(ws, title, subtitle, width):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    cell = ws.cell(row=1, column=1, value=title)
    cell.font = Font(bold=True, size=14, color=NAVY)
    cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    sub = ws.cell(row=2, column=1, value=subtitle)
    sub.font = Font(size=9, color="667085")


def _header_row(ws, row, headers, widths=None):
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=name)
        cell.font = Font(bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 26
    for col, width in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _format_date_range(ctx):
    start = ctx["range"].start_date
    end = ctx["range"].end_date
    if start == end:
        return f"Tanggal: {start.strftime('%d/%m/%Y')}"
    return f"Periode: {start.strftime('%d/%m/%Y')} – {end.strftime('%d/%m/%Y')}"


def _sheet_rekap(wb, ctx):
    ws = wb.create_sheet("Summary")
    date_display = _format_date_range(ctx)
    _title_block(ws, f"Summary — {date_display}", date_display, 4)

    _header_row(ws, 4, ["Status", "Ticket Count", "Percentage", "Notes"],
                [22, 16, 14, 44])

    total = ctx["counts"]["total"] or 1
    rows = [
        ("Open", ctx["counts"]["open"], "Pending assignment / action"),
        ("Outstanding", ctx["counts"]["outstanding"], "In progress / waiting"),
        ("Close", ctx["counts"]["closed"], "Resolved / closed"),
    ]
    row_no = 5
    for label, value, note in rows:
        ws.cell(row=row_no, column=1, value=label).font = Font(bold=True, size=10)
        ws.cell(row=row_no, column=2, value=value)
        pct = ws.cell(row=row_no, column=3, value=value / total)
        pct.number_format = "0.0%"
        ws.cell(row=row_no, column=4, value=note)
        for col in range(1, 5):
            cell = ws.cell(row=row_no, column=col)
            cell.border = BORDER
            cell.fill = PatternFill("solid", fgColor=STATUS_FILL[label])
            if col in (2, 3):
                cell.alignment = Alignment(horizontal="center")
        row_no += 1

    ws.cell(row=row_no, column=1, value="TOTAL").font = Font(bold=True, size=10)
    ws.cell(row=row_no, column=2, value=ctx["counts"]["total"]).font = Font(bold=True)
    for col in range(1, 5):
        cell = ws.cell(row=row_no, column=col)
        cell.border = BORDER
        cell.fill = PatternFill("solid", fgColor=NAVY_SOFT)
        if col == 2:
            cell.alignment = Alignment(horizontal="center")

    # Ticket type block
    row_no += 2
    ws.cell(row=row_no, column=1, value="By Ticket Type").font = Font(bold=True, size=11,
                                                                      color=NAVY)
    row_no += 1
    _header_row(ws, row_no, ["Type", "Code", "Count", ""])
    for label, code, value in [("Problem", "PR", ctx["type_counts"]["problem"]),
                               ("Feature Request", "FR", ctx["type_counts"]["feature"]),
                               ("Relocation", "RC", ctx["type_counts"]["relocation"])]:
        row_no += 1
        ws.cell(row=row_no, column=1, value=label)
        ws.cell(row=row_no, column=2, value=code).alignment = Alignment(horizontal="center")
        ws.cell(row=row_no, column=3, value=value).alignment = Alignment(horizontal="center")
        for col in range(1, 4):
            ws.cell(row=row_no, column=col).border = BORDER

    ws.freeze_panes = "A5"
    return ws


def _sheet_detail(wb, ctx):
    ws = wb.create_sheet("Ticket Details")

    from apps.tickets.models import Ticket
    has_rc = any(t.ticket_type == Ticket.Type.RELOCATION for t in ctx["tickets"])

    if has_rc:
        headers = [
            "Ticket ID", "Subject", "Status", "Level", "Type", 
            "Problem Type", "Origin Address", "New Site", "New Address",
            "Close Reason", "Agent", 
            "Created time", "Closed time", "Company"
        ]
        col_widths = [14, 42, 14, 12, 16, 20, 30, 16, 30, 50, 20, 18, 18, 25]
        wrap_cols = (2, 7, 9, 10)
        center_cols = (3, 4, 5, 8, 12, 13)
    else:
        headers = [
            "Ticket ID", "Subject", "Status", "Level", "Type", 
            "Problem Type", "Close Reason", "Agent", 
            "Created time", "Closed time", "Company"
        ]
        col_widths = [14, 42, 14, 12, 16, 20, 50, 20, 18, 18, 25]
        wrap_cols = (2, 7)
        center_cols = (3, 4, 5, 9, 10)

    date_display = _format_date_range(ctx)
    _title_block(ws, f"Ticket Details — {date_display}", date_display, len(headers))
    _header_row(ws, 4, headers, col_widths)

    for i, ticket in enumerate(ctx["tickets"], start=5):
        opened = timezone.localtime(ticket.opened_at)
        closed = timezone.localtime(ticket.closed_at) if ticket.closed_at else None
        status_label = ticket.get_status_display()
        agent_name = ticket.created_by.display_name if ticket.created_by else "—"

        if has_rc:
            values = [
                ticket.number,
                ticket.subject,
                status_label,
                ticket.get_priority_display(),
                ticket.get_ticket_type_display(),
                ticket.get_problem_category_display() if ticket.problem_category else "—",
                ticket.address or "—",
                ticket.new_site or "—",
                ticket.new_address or "—",
                ticket.close_reason or "—",
                agent_name,
                opened.strftime("%d/%m/%Y %H:%M"),
                closed.strftime("%d/%m/%Y %H:%M") if closed else "—",
                ticket.company or "—",
            ]
        else:
            values = [
                ticket.number,
                ticket.subject,
                status_label,
                ticket.get_priority_display(),
                ticket.get_ticket_type_display(),
                ticket.get_problem_category_display() if ticket.problem_category else "—",
                ticket.close_reason or "—",
                agent_name,
                opened.strftime("%d/%m/%Y %H:%M"),
                closed.strftime("%d/%m/%Y %H:%M") if closed else "—",
                ticket.company or "—",
            ]

        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.border = BORDER
            cell.font = Font(size=10)
            cell.alignment = Alignment(vertical="top",
                                       wrap_text=col in wrap_cols,
                                       horizontal="center" if col in center_cols else "left")
            if col == 3:
                cell.fill = PatternFill("solid",
                                        fgColor=STATUS_FILL.get(status_label, "FFFFFF"))

    return ws


def _sheet_agents(wb, ctx):
    ws = wb.create_sheet("Agent Performance")
    date_display = _format_date_range(ctx)
    _title_block(ws, f"Agent Performance — {date_display}", date_display, 6)

    _header_row(ws, 4, ["Agent", "Position", "Open", "Outstanding", "Close",
                        "Completion Rate"], [26, 22, 12, 16, 12, 16])

    for i, agent in enumerate(ctx["agents"], start=5):
        values = [agent["name"], agent["position"], agent["open"],
                  agent["outstanding"], agent["closed"]]
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=value)
            cell.border = BORDER
            cell.font = Font(size=10)
            if col >= 3:
                cell.alignment = Alignment(horizontal="center")
        rate = ws.cell(row=i, column=6, value=agent["rate"] / 100)
        rate.number_format = "0%"
        rate.alignment = Alignment(horizontal="center")
        rate.border = BORDER
    return ws


def build_workbook(ctx, report_type="detail"):
    """Susun workbook dari konteks laporan; kembalikan objek Workbook."""
    wb = Workbook()
    wb.remove(wb.active)

    if report_type == "rekap":
        _sheet_rekap(wb, ctx)
    elif report_type == "agent":
        _sheet_agents(wb, ctx)
    else:
        _sheet_detail(wb, ctx)

    date_display = _format_date_range(ctx)
    wb.properties.title = f"Report {date_display} — Ticketing NMS"
    wb.properties.creator = ctx.get("author", "Ticketing NMS")
    return wb


def file_name(period, date_range, report_type="detail"):
    prefix = {
        "detail": "Ticket_Details",
        "rekap": "Summary_Report",
        "agent": "Agent_Performance"
    }.get(report_type, "Report")

    d_start = date_range.start_date.strftime("%d-%m-%Y")
    d_end = date_range.end_date.strftime("%d-%m-%Y")

    if date_range.start_date == date_range.end_date:
        return f"{prefix}_{d_start}.xlsx"
    return f"{prefix}_{d_start}_sd_{d_end}.xlsx"
